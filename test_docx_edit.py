import docx

doc = docx.Document('/Users/prikshit/Downloads/Agile_Phase2_Evaluation_Template (1).docx')
print("Loaded docx successfully.")
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
